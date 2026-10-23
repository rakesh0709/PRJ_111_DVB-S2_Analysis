import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_chapter_title(doc, chapter_label, title_text):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    run1 = p.add_run(chapter_label.upper() + "\n")
    run1.font.name = "Times New Roman"
    run1.font.size = Pt(18)
    run1.font.bold = True
    run1.font.color.rgb = RGBColor(0, 0, 0)
    
    run2 = p.add_run(title_text.upper())
    run2.font.name = "Times New Roman"
    run2.font.size = Pt(18)
    run2.font.bold = True
    run2.font.color.rgb = RGBColor(0, 0, 0)
    
    # Add a thin decorative line or space
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.keep_with_next = True

def add_section_heading(doc, sec_num, sec_title):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"{sec_num} {sec_title}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

def add_subsection_heading(doc, subsec_num, subsec_title):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"{subsec_num} {subsec_title}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

def add_body_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet_p(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    if bold_prefix:
        r_b = p.add_run(bold_prefix + ": ")
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(12)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0, 0, 0)
    r_t = p.add_run(text)
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(12)
    r_t.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    set_cell_shading(cell, "F1F5F9")
    
    # Left border thick accent
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1E3A8A"/>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r_title = p.add_run(title.upper() + "\n")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(11)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)
    
    r_text = p.add_run(text)
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(10.5)
    r_text.font.color.rgb = RGBColor(15, 23, 42)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(6)

def add_custom_table(doc, caption_text, headers, data_rows, col_widths=None):
    # University standard: Table caption centered above table
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(12)
    p_cap.paragraph_format.space_after = Pt(4)
    p_cap.paragraph_format.keep_with_next = True
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(11)
    run_cap.font.bold = True
    run_cap.font.color.rgb = RGBColor(0, 0, 0)
    
    table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="94A3B8", sz="6")
    
    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(h_text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(hdr_cells[i], "1E293B") # Dark slate header
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
    
    # Repeat header row across page breaks
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    # Format Data Rows
    for r_idx, row_data in enumerate(data_rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = ""
            p = row_cells[c_idx].paragraphs[0]
            # Center if short or numeric, left-align otherwise
            val_str = str(val).strip()
            if len(val_str) < 15 or val_str.replace('.', '').replace('%', '').replace('-', '').isdigit():
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(val_str)
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0, 0, 0)
            set_cell_shading(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=80, right=80)
            
    # Apply column widths if provided
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = Inches(w)
                
    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(0)
    p_post.paragraph_format.space_after = Pt(6)

def add_custom_figure(doc, img_rel_path, caption_text, width_inches=6.0):
    abs_path = os.path.abspath(img_rel_path)
    if not os.path.exists(abs_path):
        print(f"Warning: image {abs_path} does not exist!")
        return
        
    p_img = doc.add_paragraph()
    p_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(12)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.keep_with_next = True
    p_img.add_run().add_picture(abs_path, width=Inches(width_inches))
    
    # University standard: Figure caption centered below figure
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(11)
    run_cap.font.bold = True
    run_cap.font.color.rgb = RGBColor(0, 0, 0)
