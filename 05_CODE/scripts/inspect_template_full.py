from pathlib import Path
import pptx
import os

PROJECT_ROOT = Path(__file__).resolve().parents[2]
template_path = str(PROJECT_ROOT / 'Review-2 TEMPLATE_MINI PROJECT CSE7102.pptx')
prs = pptx.Presentation(template_path)

print(f"Presentation loaded. Width: {prs.slide_width}, Height: {prs.slide_height}")
print(f"Total Slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    print(f"\n==========================================")
    print(f"SLIDE {idx+1}")
    print(f"==========================================")
    for s_idx, shape in enumerate(slide.shapes):
        print(f"\n  Shape {s_idx}: '{shape.name}' | Type: {shape.shape_type} | Position: (left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height})")
        if shape.has_table:
            print("    [TABLE DATA]")
            table = shape.table
            print(f"    Rows: {len(table.rows)}, Cols: {len(table.columns)}")
            for r_i, row in enumerate(table.rows):
                row_vals = [cell.text.replace('\n', ' ') for cell in row.cells]
                print(f"      Row {r_i}: {row_vals}")
        elif shape.has_text_frame:
            tf = shape.text_frame
            for p_i, p in enumerate(tf.paragraphs):
                text = p.text.replace('\n', '\\n')
                font_name = p.font.name if p.font else None
                font_size = p.font.size.pt if (p.font and p.font.size) else None
                bold = p.font.bold if p.font else None
                color = p.font.color.rgb if (p.font and p.font.color and hasattr(p.font.color, 'rgb')) else None
                print(f"    P{p_i}: text='{text}' | font={font_name}, size={font_size}, bold={bold}, color={color}")
