from pathlib import Path
import pptx

PROJECT_ROOT = Path(__file__).resolve().parents[2]
prs = pptx.Presentation(str(PROJECT_ROOT / 'PRJ_111_Review_2_Final.pptx'))
forbidden_terms = ['lorem', 'placeholder', 'project title', 'student name', 'roll number', 'dr./mr./ms./prof.', 'todo', 'tbd']

found_issues = []

print(f"Auditing {len(prs.slides)} slides in PRJ_111_Review_2_Final.pptx...")

for idx, slide in enumerate(prs.slides):
    slide_text = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            slide_text.append(shape.text_frame.text)
        elif shape.has_table:
            for r in shape.table.rows:
                slide_text.append(' '.join([c.text for c in r.cells]))
    
    full_text = ' '.join(slide_text).lower()
    
    for term in forbidden_terms:
        if term in full_text:
            found_issues.append((idx + 1, term, full_text[:150]))

if found_issues:
    print(f"WARNING: Found {len(found_issues)} remaining placeholders:")
    for s_num, term, snippet in found_issues:
        print(f"  Slide {s_num}: term '{term}' found in text: {snippet}")
else:
    print("SUCCESS: ZERO unresolved placeholders found across all slides!")

print("\n--- Detailed Slide Text Audit ---")
for idx, slide in enumerate(prs.slides):
    slide_text = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            slide_text.append(shape.text_frame.text.strip().replace('\n', ' '))
    print(f"Slide {idx+1:02d} Title/Content: {' | '.join(slide_text[:2])}")
